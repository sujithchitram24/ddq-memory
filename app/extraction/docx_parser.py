"""DOCX extraction with python-docx.

A .docx is a ZIP of XML. Paragraphs and tables are siblings in the document body, so we
walk the body in order to keep reading order. Merged cells appear in python-docx as the
same underlying cell repeated, which we use to detect two-row headers.
"""
import io
import zipfile

import docx
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph

from app.extraction import errors
from app.extraction.errors import ExtractionError
from app.extraction.ir import ExtractedDocument, Paragraph, Table
from app.extraction.pdf_parser import attach_titles_and_footnotes
from app.extraction.tables import clean, combine_header_rows


def _has_horizontal_merge(row) -> bool:
    tcs = [c._tc for c in row.cells]
    return any(a is b for a, b in zip(tcs, tcs[1:]))


def _table(t: DocxTable, index: int, heading: str | None) -> Table:
    rows = list(t.rows)
    grid = [[clean(c.text) for c in r.cells] for r in rows]
    header_rows = 2 if len(rows) > 2 and _has_horizontal_merge(rows[0]) else 1
    if header_rows == 2:
        header = combine_header_rows(grid[0], grid[1])
        flags = ["multi_row_header"]
    else:
        header = [h or "" for h in grid[0]]
        flags = []
    return Table(
        header=header,
        rows=grid[header_rows:],
        location={"table_index": index, "section": heading},
        quality_flags=flags,
    )


def parse_docx(data: bytes) -> ExtractedDocument:
    try:
        d = docx.Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ExtractionError(errors.CORRUPT_FILE, f"DOCX could not be opened: {type(exc).__name__}") from exc

    blocks: list[Paragraph | Table] = []
    heading: str | None = None
    table_index = 0
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = DocxParagraph(child, d)
            text = clean(p.text)
            if not text:
                continue
            if p.style is not None and p.style.name.startswith("Heading"):
                heading = text
            blocks.append(Paragraph(text=text, location={"block": len(blocks), "section": heading}))
        elif tag == "tbl":
            blocks.append(_table(DocxTable(child, d), table_index, heading))
            table_index += 1

    if not blocks:
        raise ExtractionError(errors.EMPTY_DOCUMENT, "DOCX contains no text or tables")
    return ExtractedDocument(file_format="docx", parser="python-docx", blocks=attach_titles_and_footnotes(blocks))
