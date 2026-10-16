"""XLSX extraction with openpyxl.

Things that go wrong with spreadsheets, and what we do:
- Numbers vs display: 0.0125 formatted as "0.00%" is shown as 1.25%. We render using
  the cell's number_format so the text matches what a human saw.
- Formulas: openpyxl does not calculate. With data_only=True we get the value Excel
  cached on last save; a file written by a script has no cache -> None. We flag it.
- Merged cells: only the top-left cell holds the value. A merge spanning the whole row
  is treated as a title; other merges are filled with the top-left value.
- Multiple sheets: each sheet becomes its own table (or paragraphs, if one column).
"""
import io
import re
import zipfile
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.extraction import errors
from app.extraction.errors import ExtractionError
from app.extraction.ir import ExtractedDocument, Paragraph, Table
from app.extraction.tables import FOOTNOTE_LINE, clean, finalize

PCT_DECIMALS = re.compile(r"0\.(0+)%")


def render(value: object, number_format: str) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)) and "%" in number_format:
        m = PCT_DECIMALS.search(number_format)
        decimals = len(m.group(1)) if m else 0
        return f"{value * 100:.{decimals}f}%"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.isoformat()[:10]
    return clean(value)


def _looks_like_data(cells: list[str | None]) -> bool:
    from app.processing.values import parse_value

    return any(c and parse_value(c).kind in ("percent", "money") or (c and re.fullmatch(r"-?\d+\.\d+", c)) for c in cells)


def parse_xlsx(data: bytes) -> ExtractedDocument:
    try:
        wb_values = load_workbook(io.BytesIO(data), data_only=True)
        wb_formulas = load_workbook(io.BytesIO(data), data_only=False)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise ExtractionError(errors.CORRUPT_FILE, f"XLSX could not be opened: {type(exc).__name__}") from exc

    blocks: list[Paragraph | Table] = []
    warnings: list[str] = []
    table_index = 0
    for ws in wb_values.worksheets:
        if ws.sheet_state != "visible":
            warnings.append(f"sheet '{ws.title}' is hidden; skipped")
            continue
        wf = wb_formulas[ws.title]
        max_r, max_c = ws.max_row, ws.max_column
        grid: dict[tuple[int, int], str | None] = {}
        formula_cells: list[str] = []
        for r in range(1, max_r + 1):
            for c in range(1, max_c + 1):
                cell = ws.cell(r, c)
                grid[(r, c)] = render(cell.value, cell.number_format)
                fval = wf.cell(r, c).value
                if cell.value is None and isinstance(fval, str) and fval.startswith("="):
                    formula_cells.append(f"{get_column_letter(c)}{r}")

        title_rows: set[int] = set()
        for rng in ws.merged_cells.ranges:
            top_left = grid[(rng.min_row, rng.min_col)]
            if rng.min_row == rng.max_row and rng.min_col == 1 and rng.max_col >= max_c:
                title_rows.add(rng.min_row)
                continue
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    grid[(r, c)] = top_left

        used_cols = [c for c in range(1, max_c + 1) if any(grid[(r, c)] for r in range(1, max_r + 1))]
        rows = [r for r in range(1, max_r + 1) if any(grid[(r, c)] for c in used_cols)]
        if not rows:
            continue

        # Leading rows with a single filled cell are titles/banners -> paragraphs.
        title: str | None = None
        while rows and (rows[0] in title_rows or sum(1 for c in used_cols if grid[(rows[0], c)]) == 1) and len(used_cols) > 1:
            text = next(grid[(rows[0], c)] for c in used_cols if grid[(rows[0], c)])
            blocks.append(Paragraph(text=text, location={"sheet": ws.title, "cell": f"A{rows[0]}"}))
            title = text
            rows = rows[1:]

        if len(used_cols) <= 1:
            for r in rows:
                text = grid[(r, used_cols[0])]
                blocks.append(Paragraph(text=text, location={"sheet": ws.title, "cell": f"{get_column_letter(used_cols[0])}{r}"}))
            continue
        if not rows:
            continue

        flags: list[str] = []
        header_vals = [grid[(rows[0], c)] for c in used_cols]
        if len(rows) == 1:
            header, data_rows, header_row = ["Field", "Value"][: len(used_cols)], rows, None
            flags.append("synthesized_header")
        elif all(v is None for v in header_vals) or _looks_like_data(header_vals):
            # A blank header row is skipped as empty, so the first DATA row would silently
            # become the header. Values like "1.5%" in a header mean exactly that.
            header_row = None if _looks_like_data(header_vals) else rows[0]
            data_rows = rows if header_row is None else rows[1:]
            header = [f"column_{i+1}" for i in range(len(used_cols))]
            flags.append("missing_header")
        else:
            header, data_rows, header_row = [v or "" for v in header_vals], rows[1:], rows[0]
        if formula_cells:
            flags.append("formula_without_cached_value")
            warnings.append(f"sheet '{ws.title}': formulas without cached values at {', '.join(formula_cells)}")

        table = Table(
            title=title or ws.title,
            header=header,
            rows=[[grid[(r, c)] for c in used_cols] for r in data_rows],
            location={
                "sheet": ws.title,
                "table_index": table_index,
                "header_row": header_row,
                "row_numbers": data_rows,
                "columns": [get_column_letter(c) for c in used_cols],
            },
            quality_flags=flags,
        )
        blocks.append(finalize(table))
        table_index += 1

    # Notes-style sheets: footnotes such as "(1) ..." become paragraphs; attach them to
    # any table whose cells reference the same marker.
    notes = [b for b in blocks if isinstance(b, Paragraph) and FOOTNOTE_LINE.match(b.text)]
    for t in (b for b in blocks if isinstance(b, Table)):
        cells = " ".join(v for row in t.rows for v in row if v)
        for n in notes:
            marker = n.text.split(")", 1)[0] + ")"
            if marker in cells:
                t.footnotes.append(n.text)

    if not blocks:
        raise ExtractionError(errors.EMPTY_DOCUMENT, "XLSX has no visible content")
    return ExtractedDocument(file_format="xlsx", parser="openpyxl", blocks=blocks, warnings=warnings)
